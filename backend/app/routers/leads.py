import csv
import io
import redis.asyncio as aioredis
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, desc
from app.database import get_db
from app.models.lead import Lead, LeadStatus, LeadSource
from app.config import get_settings
from app.schemas.lead import LeadResponse, LeadUpdate, LeadSearchRequest
from app.services.leads_service import LeadsService
from app.services.whatsapp_service import WhatsAppService

router = APIRouter(prefix="/leads", tags=["leads"])


@router.get("")
async def get_leads(
    page: int = 1,
    limit: int = 20,
    status: str = None,
    db: AsyncSession = Depends(get_db),
):
    query = select(Lead).order_by(desc(Lead.created_at))
    if status:
        query = query.where(Lead.status == status)
    total = await db.scalar(select(func.count()).select_from(Lead))
    result = await db.execute(query.offset((page - 1) * limit).limit(limit))
    leads = result.scalars().all()
    return {
        "leads": [LeadResponse.model_validate(l) for l in leads],
        "total": total,
        "page": page,
    }


@router.post("/search")
async def search_leads(req: LeadSearchRequest, db: AsyncSession = Depends(get_db)):
    # Clear previous uncontacted Maps leads so each search starts fresh
    prev = await db.execute(
        select(Lead).where(Lead.source == LeadSource.google_maps, Lead.status == LeadStatus.new)
    )
    old_leads = prev.scalars().all()
    if old_leads:
        r = aioredis.from_url(get_settings().REDIS_URL, decode_responses=True)
        for lead in old_leads:
            await r.delete(f"lead:phone:{lead.phone}")
            await db.delete(lead)
        await db.commit()

    svc = LeadsService()
    try:
        places = await svc.search_places(req.query, req.location, req.radius_km)
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Google Places error: {e}")

    created = 0
    for place in places:
        lead = await svc.process_lead(place, db)
        if lead:
            created += 1

    return {"found": len(places), "new_leads": created}


@router.post("/import")
async def import_contacts(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    content = await file.read()
    filename = (file.filename or "").lower()
    contacts: list[dict] = []

    if filename.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            ws = wb.active
            headers = [str(c.value or "").strip().lower() for c in next(ws.iter_rows(max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = dict(zip(headers, row))
                phone = str(row_dict.get("phone") or row_dict.get("telefono") or row_dict.get("celular") or "")
                name  = str(row_dict.get("name")  or row_dict.get("nombre")   or row_dict.get("negocio") or "")
                contacts.append({"phone": phone.strip(), "name": name.strip()})
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error al leer Excel: {e}")
    elif filename.endswith(".csv") or "text" in (file.content_type or ""):
        try:
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            for row in reader:
                lower = {k.lower().strip(): v for k, v in row.items()}
                phone = lower.get("phone") or lower.get("telefono") or lower.get("celular") or ""
                name  = lower.get("name")  or lower.get("nombre")   or lower.get("negocio") or ""
                contacts.append({"phone": str(phone).strip(), "name": str(name).strip()})
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error al leer CSV: {e}")
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Usá .csv o .xlsx")

    wa = WhatsAppService()
    imported = skipped = errors = 0

    for c in contacts:
        phone = "".join(ch for ch in c.get("phone", "") if ch.isdigit())
        if not phone:
            skipped += 1
            continue
        try:
            existing = await db.scalar(select(Lead).where(Lead.phone == phone))
            if existing:
                skipped += 1
                continue
            lead = Lead(
                phone=phone,
                name=c.get("name", ""),
                business_name=c.get("name", ""),
                source=LeadSource.google_maps,
                status=LeadStatus.new,
            )
            db.add(lead)
            await db.commit()
            await db.refresh(lead)
            try:
                await wa.send_template(
                    phone=phone,
                    template_name="primer_contacto_maps",
                    language="es",
                    params=[lead.business_name or lead.name or ""],
                )
                lead.status = LeadStatus.contacted
                await db.commit()
            except Exception:
                pass
            imported += 1
        except Exception:
            errors += 1

    return {"imported": imported, "skipped": skipped, "errors": errors, "total": len(contacts)}


@router.get("/stats")
async def get_lead_stats(db: AsyncSession = Depends(get_db)):
    stats = {}
    for status in LeadStatus:
        count = await db.scalar(select(func.count()).select_from(Lead).where(Lead.status == status))
        stats[status.value] = count or 0
    return stats


@router.get("/{lead_id}", response_model=LeadResponse)
async def get_lead(lead_id: int, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Lead).where(Lead.id == lead_id))
    lead = result.scalar_one_or_none()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    return lead


@router.put("/{lead_id}/status", response_model=LeadResponse)
async def update_lead(lead_id: int, data: LeadUpdate, db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Lead).where(Lead.id == lead_id))
    lead = result.scalar_one_or_none()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    if data.status:
        lead.status = data.status
    if data.notes:
        lead.notes = data.notes
    if data.kommo_lead_id:
        lead.kommo_lead_id = data.kommo_lead_id
    await db.commit()
    await db.refresh(lead)
    return lead


@router.post("/{lead_id}/contact")
async def contact_lead(lead_id: int, db: AsyncSession = Depends(get_db)):
    """Mark lead as contacted and (re)create in Kommo Maps pipeline if missing."""
    result = await db.execute(select(Lead).where(Lead.id == lead_id))
    lead = result.scalar_one_or_none()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")

    from app.services.kommo_service import KommoService, STATUS_MAPS_CONTACTADO, PIPELINE_MAPS

    kommo_error = None
    if not lead.kommo_lead_id:
        try:
            kommo = KommoService()
            kommo_lead = await kommo.create_lead_from_maps(
                phone=lead.phone,
                business_name=lead.business_name or lead.name or "",
                address=lead.address or "",
            )
            if kommo_lead.get("id"):
                lead.kommo_lead_id = str(kommo_lead["id"])
        except Exception as e:
            kommo_error = str(e)
            print(f"[contact_lead] kommo error: {e}")
    else:
        try:
            kommo = KommoService()
            await kommo.move_lead_stage(int(lead.kommo_lead_id), STATUS_MAPS_CONTACTADO)
        except Exception as e:
            print(f"[contact_lead] kommo move error: {e}")

    lead.status = LeadStatus.contacted
    await db.commit()
    await db.refresh(lead)

    wa_link = f"https://wa.me/{lead.phone}"
    return {
        "lead": LeadResponse.model_validate(lead),
        "wa_link": wa_link,
        "kommo_synced": lead.kommo_lead_id is not None,
        "kommo_error": kommo_error,
    }
